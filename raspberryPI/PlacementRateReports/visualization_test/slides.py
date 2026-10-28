"""Rendered dashboard pages -> weekly PowerPoint deck (title, overall, programs)."""
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

TEMPLATE = Path(__file__).parent / 'assets/weekly_slides_template.pptx'
MARGIN = Emu(320000)
TITLE_FONT = 'Calibri'


def ordinal_suffix(day):
    return 'th' if 10 <= day % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(day % 10, 'th')


def deck_filename(report_date):
    """October_9th_Placement_Rate_Visualization.pptx, matching the manual naming."""
    day = report_date.day
    return f'{report_date:%B}_{day}{ordinal_suffix(day)}_Placement_Rate_Visualization.pptx'


def _drop_slide(presentation, index):
    ids = presentation.slides._sldIdLst
    slide_id = list(ids)[index]
    presentation.part.drop_rel(slide_id.rId)
    ids.remove(slide_id)


def _title_slide(slide, report_date):
    box = slide.shapes.add_textbox(Emu(560000), Emu(2420000), Emu(4200000), Emu(2000000))
    box.name = 'Report Date Title'
    frame = box.text_frame
    frame.word_wrap = True
    lines = [(f'{report_date:%B} {report_date.day}', ordinal_suffix(report_date.day)),
             ('Placement', None), ('Overview', None)]
    for i, (text, superscript) in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.line_spacing = 0.9
        for part, raised in ((text, False), (superscript, True)):
            if not part:
                continue
            run = paragraph.add_run()
            run.text = part
            run.font.name = TITLE_FONT
            run.font.size = Pt(44)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            if raised:
                run.font._rPr.set('baseline', '30000')


def _picture_slide(presentation, prototype, image):
    slide = presentation.slides.add_slide(prototype.slide_layout)
    background = prototype._element.cSld.bg
    if background is not None:
        slide._element.cSld.insert(0, deepcopy(background))
    for shape in prototype.shapes:
        slide.shapes._spTree.append(deepcopy(shape._element))
    with Image.open(image) as img:
        ratio = img.width / img.height
    width, height = presentation.slide_width - 2 * MARGIN, presentation.slide_height - 2 * MARGIN
    if width / height > ratio:
        width = int(height * ratio)
    else:
        height = int(width / ratio)
    left = (presentation.slide_width - width) // 2
    top = (presentation.slide_height - height) // 2
    slide.shapes.add_picture(str(image), left, top, width, height).name = Path(image).stem


def build_deck(images, output, *, report_date, template=TEMPLATE):
    """images: rendered dashboard PNGs in slide order (overall first, then programs)."""
    images = [Path(i) for i in images]
    if not images or any(not i.is_file() for i in images):
        raise ValueError('Slide deck requires every rendered dashboard image')
    presentation = Presentation(str(template))
    if len(presentation.slides) != 2:
        raise ValueError('Slide template must contain exactly a title slide and a picture prototype')
    title, prototype = presentation.slides
    _title_slide(title, report_date)
    for image in images:
        _picture_slide(presentation, prototype, image)
    _drop_slide(presentation, 1)
    presentation.core_properties.title = f'{report_date:%B} {report_date.day}{ordinal_suffix(report_date.day)} Placement Overview'
    output = Path(output)
    temporary = output.with_suffix('.pptx.tmp')
    try:
        presentation.save(str(temporary))
        if len(Presentation(str(temporary)).slides) != len(images) + 1:
            raise RuntimeError('Slide deck page count mismatch')
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return {'slides': len(images) + 1, 'bytes': output.stat().st_size}
