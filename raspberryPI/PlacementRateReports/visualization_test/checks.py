"""Cross-check a built dashboard against the report workbook it was built from.

Run before screenshots, so a dashboard that disagrees with its report is never sent:
  - the weekly chart data (NittyGrittySheet) has exactly the report's Fridays in the
    class window, and every point's % placed and status counts equal the report's;
  - each page names the class (title, table header N24, weekly chart titles);
  - this month's cell in the comparison table equals the report-date % placed;
  - every chart's cached series has one point per week, so Excel and LibreOffice draw
    the same data without recalculating.
"""
from openpyxl.utils import get_column_letter
from Visualization.update_dashboard import read_history, monthly_row, Package, PAGES, NITTY, NS, STATUSES


def check_dashboard(dashboard, source, *, pages, class_year, as_of, start_date):
    errors = []
    history = read_history(source, pages, as_of, start_date)
    fridays = sorted({d for records in history.values() for d in records if d.weekday() == 4})
    package = Package(dashboard)
    for page in pages:
        index = PAGES.index(page)
        rate_row, status_header = 24 + 17 * index, 26 + 17 * index
        status_rows = list(range(status_header + 1, status_header + 12))
        dates = []
        for col in range(2, 2 + len(fridays) + 1):
            label = package.get(NITTY, f'{get_column_letter(col)}{rate_row - 1}')
            if label is None:
                break
            dates.append((col, label))
        expected = [d.strftime('%m/%d/%Y') for d in fridays]
        if [label for _, label in dates] != expected:
            errors.append(f'{page}: weekly dates {[l for _, l in dates][-3:]} != report Fridays {expected[-3:]} '
                          f'({len(dates)} vs {len(expected)})')
            continue
        for (col, label), when in zip(dates, fridays):
            letter = get_column_letter(col)
            source_point = history[page].get(when)
            if source_point is None:
                errors.append(f'{page}: report has no {label} observation for a charted week')
                continue
            charted = [package.get(NITTY, f'{letter}{row}') for row in [rate_row] + status_rows]
            labels = [package.get(NITTY, f'A{row}') for row in status_rows]
            by_label = dict(zip(labels, charted[1:]))
            counts = [by_label.get(status) for status in STATUSES]
            if charted[0] != source_point['rate'] or counts != source_point['counts']:
                errors.append(f'{page} {label}: chart {charted[0]!r} {counts} != report '
                              f'{source_point["rate"]!r} {source_point["counts"]}')
        if str(package.get(page, 'N24')) != str(class_year):
            errors.append(f'{page}: comparison table current class is {package.get(page, "N24")!r}, expected {class_year}')
        if str(class_year) not in str(package.get(page, 'B1') or ''):
            errors.append(f'{page}: title {package.get(page, "B1")!r} does not name the Class of {class_year}')
        row = monthly_row(class_year, as_of)
        current = history[page].get(as_of)
        if row is not None and current is not None and package.get(page, f'N{row}') != current['rate']:
            errors.append(f'{page}: this month (N{row}) shows {package.get(page, f"N{row}")!r}, '
                          f'report says {current["rate"]!r} on {as_of:%m/%d/%Y}')
    for part in package.chart_parts():
        root = package.xml(part)
        for title in root.findall('.//c:title', NS):
            text = ''.join(t.text or '' for t in title.findall('.//a:t', NS))
            if text.lower().startswith('weekly') and str(class_year) not in text:
                errors.append(f'{part}: chart title {text!r} does not name the Class of {class_year}')
        for ref in root.findall('.//c:numRef', NS) + root.findall('.//c:strRef', NS):
            formula = ref.find('c:f', NS)
            if formula is None or not (formula.text or '').startswith(NITTY + '!') or ':' not in formula.text:
                continue
            count = ref.find('.//c:ptCount', NS)
            if count is None or int(count.get('val')) != len(fridays):
                errors.append(f'{part}: chart series {formula.text} caches '
                              f'{count.get("val") if count is not None else "no"} points, expected {len(fridays)}')
    return {'errors': errors, 'weeks': len(fridays), 'pages': list(pages)}
