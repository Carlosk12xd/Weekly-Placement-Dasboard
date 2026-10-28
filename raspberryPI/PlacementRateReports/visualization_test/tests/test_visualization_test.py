"""Visualization test pipeline; offline fixtures, rendering mocked, SMTP never called.

The supplied fixtures end on 2026-08-21, before the Class of 2027 becomes the report class
in October, so setUpClass adds an October 9, 2026 week with the real report builders and
mocked SQL (tests/october_fixture.py).

Run from PlacementRateReports:  python -m unittest discover -s visualization_test/tests -t . -v
"""
from datetime import date
from email import policy
from email.parser import BytesParser
import json
import os
from pathlib import Path
import shutil
import tempfile
from unittest import TestCase, main
from unittest.mock import patch
import yaml

from core.daily_runner import load_cohorts
from core.report_artifacts import ROOT
from tests.october_fixture import build_week
from Visualization.update_dashboard import Package, OVERALL, NITTY, NS, read_history
from visualization_test import run as runner
from visualization_test.checks import check_dashboard
from visualization_test.dashboard import (academic_class_year, table_years, build_dashboard,
    write_table, template_seed)
from visualization_test.emails import director_greeting, subject_line, body_text

FIXTURES = ROOT / 'tests/fixtures'
LEADERSHIP = 'LeadershipReport/Reports/2027_weekly_placement_report.xlsx'
FINANCE = 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
REPORT26 = FIXTURES / 'LeadershipReport/Reports/2026_weekly_placement_report.xlsx'
DAY = date(2026, 10, 9)
START = date(2026, 6, 22)
TEMPLATE_2026_08 = Path(__file__).parent / 'fixtures/placement_dashboard_template_2026-08.xlsx'


class VisualizationTestPipeline(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.work = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.work.name)
        # Copy so an open Excel window on the template does not block the tests. The
        # default is the August 2026 template, so a newer live template (with later
        # months recorded) does not change the results.
        cls.template = cls.folder / 'template.xlsx'
        shutil.copyfile(os.environ.get('VT_TEST_TEMPLATE', TEMPLATE_2026_08), cls.template)
        # A config with no test inboxes, independent of the live config.yml.
        cls.config = cls.folder / 'no-recipients.yml'
        cls.config.write_text(yaml.safe_dump({**yaml.safe_load((runner.HERE / 'config.yml').read_text()),
                                              'test_recipients': [], 'box_upload_addresses': {}}))
        cls.source = cls.folder / 'source'
        shutil.copytree(FIXTURES, cls.source)
        cohort = next(c for c in load_cohorts()[1] if c.id == '2027')
        build_week(cls.source / LEADERSHIP, cls.source / LEADERSHIP, cohort=cohort, report_date=DAY)
        build_week(cls.source / FINANCE, cls.source / FINANCE, cohort=cohort, report_date=DAY, programs=('BSFin',))
        cls.report27 = cls.source / FINANCE
        cls.output = cls.folder / 'BSFin.xlsx'
        cls.meta = build_dashboard(cls.report27, cls.output, class_year=2027, as_of=DAY,
                                   start_date=START, programs=('BSFin',),
                                   template=cls.template, state_path=cls.folder / 'state.json',
                                   prior_source=REPORT26)

    @classmethod
    def tearDownClass(cls):
        cls.work.cleanup()

    def test_academic_year(self):
        # The report class switches on October 1 (core.academic_year).
        self.assertEqual(academic_class_year(date(2026, 9, 25)), 2026)
        self.assertEqual(academic_class_year(date(2026, 10, 2)), 2027)
        self.assertEqual(academic_class_year(date(2027, 9, 24)), 2027)
        self.assertEqual(table_years(2027), [2027, 2026, 2025, 2024, 2023, 2022])

    def test_email_text(self):
        self.assertEqual(director_greeting(['Bob', 'Audrey']), 'Audrey & Bob')
        self.assertEqual(director_greeting(['Tanya', 'amy', 'Bob']), 'amy, Bob & Tanya')
        self.assertEqual(subject_line([2027], 'BSFin', DAY), 'Class of 2027: BSFin Weekly Placement Report 2026-10-09')
        self.assertEqual(body_text(['Tanya', 'Amy'], date(2026, 10, 16)),
                         'Amy & Tanya,\n\nHere is the weekly placement visualization for 10-16. '
                         'Please let us know if you have any questions.\n\nBest,\nBCC Data Team')

    def test_table_rolls_in_october_with_both_classes_from_their_own_reports(self):
        package = Package(self.output)
        for page in (OVERALL, 'BSFin'):
            self.assertEqual([str(int(float(package.get(page, f'{c}24')))) for c in 'NOPQRS'],
                             ['2027', '2026', '2025', '2024', '2023', '2022'])
        for page in package.sheets:
            if page != NITTY:  # Numeric table headers make Excel repair the workbook.
                self.assertTrue(all(package.cell(page, f'{c}24').get('t') == 'inlineStr' for c in 'NOPQRS'), page)
        current = read_history(self.report27, (OVERALL, 'BSFin'), DAY, START)
        prior = read_history(REPORT26, (OVERALL, 'BSFin'), DAY, date(2025, 1, 1), require_current=False)
        for page in (OVERALL, 'BSFin'):
            # Class of 2027: Aug 2026 (its last August week), Oct 2026 = this report.
            self.assertEqual(package.get(page, 'N25'), current[page][date(2026, 8, 21)]['rate'])
            self.assertEqual(package.get(page, 'N27'), current[page][DAY]['rate'])
            self.assertIsNone(package.get(page, 'N28'))                               # Nov 2026 not yet
            # Class of 2026 from its own final report: August 2026 = its 08/31 observation.
            self.assertEqual(package.get(page, 'O37'), prior[page][date(2026, 8, 31)]['rate'])
        self.assertTrue(self.meta['prior_class_updated'])

    def test_titles_and_visible_tabs(self):
        package = Package(self.output)
        titles = {''.join(t.text or '' for t in title.findall('.//a:t', NS))
                  for part in package.chart_parts() for title in package.xml(part).findall('.//c:title', NS)}
        weekly = {t for t in titles if t.lower().startswith('weekly')}
        self.assertEqual(weekly, {'Weekly Class of 2027 Placement Trend', 'Weekly Class of 2027 Search Status'})
        from openpyxl import load_workbook
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb.active.title, OVERALL)
            # Director workbook: only the MSB tab and this director's programs are visible.
            self.assertEqual([ws.title for ws in wb if ws.sheet_state == 'visible'], [OVERALL, 'BSFin'])
        finally:
            wb.close()

    def test_dashboard_matches_its_report(self):
        result = check_dashboard(self.output, self.report27, pages=self.meta['pages'], class_year=2027,
                                 as_of=DAY, start_date=START)
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['weeks'], len(self.meta['weekly_dates']))
        self.assertEqual(self.meta['weekly_dates'][0], '2026-06-26')  # history starts the last week of June

    def test_next_academic_year_rolls_saved_class_into_history(self):
        state = json.loads((self.folder / 'state.json').read_text())
        package = Package(self.output)
        write_table(package, ('BSFin',), 2028, {}, state, template_seed(self.template), date(2027, 10, 8))
        self.assertEqual([str(int(float(package.get('BSFin', f'{c}24')))) for c in 'NOPQRS'],
                         ['2028', '2027', '2026', '2025', '2024', '2023'])
        self.assertEqual(package.get('BSFin', 'O27'), state['BSFin']['2027']['27']['rate'])
        self.assertEqual(package.get('BSFin', 'P37'), state['BSFin']['2026']['37']['rate'])
        self.assertIsNone(package.get('BSFin', 'N25'))

    def test_saved_values_after_report_date_are_not_used(self):
        state = {'BSFin': {'2027': {'28': {'rate': 0.5, 'date': '2026-11-06'}}}}
        package = Package(self.output)
        write_table(package, ('BSFin',), 2027, {}, state, {}, DAY)
        self.assertIsNone(package.get('BSFin', 'N28'))

    @staticmethod
    def fake_images(_, pages, folder, **kwargs):
        from PIL import Image
        Path(folder).mkdir(parents=True, exist_ok=True)
        shots = [Path(folder) / f'{i:02d}-{p}.png' for i, p in enumerate(pages, 1)]
        for shot in shots:
            Image.new('RGB', (400, 225), 'white').save(shot)
        return shots

    def test_run_previews_test_only_email(self):
        run_root = self.folder / 'runs'
        args = ['--date', str(DAY), '--audience', 'Tanya', '--source-root', str(self.source),
                '--run-root', str(run_root), '--template', str(self.template), '--config', str(self.config)]
        with patch.object(runner, 'render_sheet_images', side_effect=self.fake_images), \
             patch('smtplib.SMTP', side_effect=AssertionError('SMTP must not be called')):
            self.assertEqual(runner.main(args), 0)
            with self.assertRaises(SystemExit):
                runner.main(args + ['--send'])  # No test_recipients configured.
        folder = run_root / str(DAY) / 'Tanya'
        msg = BytesParser(policy=policy.default).parsebytes((folder / 'email.eml').read_bytes())
        self.assertEqual(msg['Subject'], 'Class of 2027: BSFin Weekly Placement Report 2026-10-09')
        self.assertNotIn('tharmon@byu.edu', str(msg['To']))
        self.assertIn('tharmon@byu.edu', msg['X-Intended-To'])
        self.assertEqual([a.get_filename() for a in msg.iter_attachments()],
                         ['2027-Dashboard-BSFin-2026-10-09.xlsx', 'October_9th_Placement_Rate_Visualization.pptx'])
        self.assertTrue(msg.get_body().get_content().replace('\r\n', '\n').startswith('Amy & Tanya,\n\n'))
        manifest = json.loads((folder / 'manifest.json').read_text())
        self.assertEqual(manifest['routes']['email']['envelope'], [])
        self.assertNotIn('box', manifest['routes'])
        self.assertEqual(manifest['source'], '2027-WeeklyPlacement-BSFin.xlsx')
        self.assertEqual(manifest['checks']['errors'], 0)
        self.assertEqual(json.loads((folder / 'checks.json').read_text())['errors'], [])

    def test_stale_report_is_refused_before_any_slide(self):
        run_root = self.folder / 'stale'
        with patch.object(runner, 'render_sheet_images', side_effect=AssertionError('must not render')):
            code = runner.main(['--date', '2026-10-16', '--audience', 'Tanya', '--source-root', str(self.source),
                                '--run-root', str(run_root), '--template', str(self.template),
                                '--config', str(self.config)])
        self.assertEqual(code, 1)
        summary = json.loads((run_root / '2026-10-16/summary.json').read_text())
        self.assertIn('no observation for 2026-10-16', summary['results'][0]['error'])

    def test_director_deck_from_leadership_report(self):
        run_root = self.folder / 'leadership-source'
        with patch.object(runner, 'render_sheet_images', side_effect=self.fake_images) as render, \
             patch('smtplib.SMTP', side_effect=AssertionError('SMTP must not be called')):
            self.assertEqual(runner.main(['--date', str(DAY), '--audience', 'Bob', '--source-root', str(self.source),
                                          '--run-root', str(run_root), '--template', str(self.template),
                                          '--config', str(self.config), '--director-source', 'leadership']), 0)
        self.assertEqual(list(render.call_args.args[1]), [OVERALL, 'BSIS', 'MISM'])
        folder = run_root / str(DAY) / 'Bob'
        manifest = json.loads((folder / 'manifest.json').read_text())
        self.assertEqual(manifest['source'], '2027_weekly_placement_report.xlsx')
        from pptx import Presentation
        self.assertEqual(len(Presentation(folder / manifest['deck']).slides), 4)  # Title, MSB, BSIS, MISM

    def test_box_upload_route_has_workbook_and_labelled_deck(self):
        config = yaml.safe_load((runner.HERE / 'config.yml').read_text())
        config['box_upload_addresses'] = {'2027': ['test_folder.abc123@u.box.com'], '2026': ['old.x1@u.box.com']}
        config_path = self.folder / 'config.yml'
        config_path.write_text(yaml.safe_dump(config))
        run_root = self.folder / 'box-runs'
        with patch.object(runner, 'render_sheet_images', side_effect=self.fake_images), \
             patch('smtplib.SMTP', side_effect=AssertionError('SMTP must not be called')):
            self.assertEqual(runner.main(['--date', str(DAY), '--audience', 'Tanya', '--config', str(config_path),
                                          '--source-root', str(self.source), '--run-root', str(run_root),
                                          '--template', str(self.template)]), 0)
        folder = run_root / str(DAY) / 'Tanya'
        manifest = json.loads((folder / 'manifest.json').read_text())
        self.assertEqual(manifest['routes']['box']['envelope'], ['test_folder.abc123@u.box.com'])
        msg = BytesParser(policy=policy.default).parsebytes((folder / 'box.eml').read_bytes())
        self.assertEqual([a.get_filename() for a in msg.iter_attachments()],
                         ['2027-Dashboard-BSFin-2026-10-09.xlsx',
                          'October_9th_Placement_Rate_Visualization-BSFin.pptx'])


if __name__ == '__main__':
    main()
