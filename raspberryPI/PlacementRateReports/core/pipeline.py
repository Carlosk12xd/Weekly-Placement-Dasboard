"""Completed report snapshot -> dashboard/PDF -> journaled email routes."""
from datetime import date, datetime, timezone
from pathlib import Path
import hashlib
import json
import logging
from .report_artifacts import ROOT, atomic_copy, sha256, write_json, verify_file
from .delivery import env_list, message, save_route, deliver
from Visualization.build_visualization import build_visualization, load_profile
from Visualization.render_excel import render_pdf
from Visualization.update_dashboard import read_history, PAGES, OVERALL

logger = logging.getLogger(__name__)


def namespace_for(options):
    namespace = 'preview' if options.dry_run else 'delivery'
    if options.test_recipient:
        namespace += '-pilot-' + hashlib.sha256(options.test_recipient.encode()).hexdigest()[:12]
    return namespace


def execute_bundle(cohort, mode, report_date, audience, rolling, programs, recipients, options):
    result = {'cohort': str(cohort.id), 'audience': audience, 'status': 'failed'}
    if not options.offline and not options.retry and report_date != date.today():
        return {**result, 'error': 'Live builds require today; use offline fixtures or a saved retry'}
    directory = Path(options.run_root) / namespace_for(options) / str(report_date) / str(cohort.id) / audience
    result['directory'] = str(directory)
    manifest_path = directory / 'manifest.json'
    identity = {'cohort': str(cohort.id), 'report_date': str(report_date), 'mode': mode,
                'audience': audience, 'programs': list(programs) if programs is not None else None,
                'offline': options.offline}
    manifest = {'identity': identity, 'created_at': datetime.now(timezone.utc).isoformat(),
                'artifacts': {}, 'routes': {}}
    # A corrupt or mismatched ledger is evidence, not a blank run to overwrite.
    try:
        directory.mkdir(parents=True, exist_ok=True)
        if manifest_path.exists():
            manifest = json.loads(manifest_path.read_text())
            if manifest['identity'] != identity:
                raise RuntimeError('Saved run identity differs; inspect the existing run')
            if not isinstance(manifest['artifacts'], dict) or not isinstance(manifest['routes'], dict):
                raise ValueError('Invalid saved artifact or route ledger')
        elif options.retry:
            raise RuntimeError('No saved run to retry')
    except Exception as exc:
        return {**result, 'error': f'{type(exc).__name__}: {exc}'}
    try:
        prepare_and_deliver(cohort, mode, report_date, audience, Path(rolling), programs,
                            recipients, options, directory, manifest)
        manifest['status'] = 'failed' if manifest.get('visualization_error') else ('preview' if options.dry_run else 'completed')
        manifest.pop('error', None)
    except Exception as exc:
        manifest['status'] = 'failed'
        manifest['error'] = f'{type(exc).__name__}: {exc}'
        logger.error('%s/%s failed: %s', cohort.id, audience, manifest['error'])
    try:
        write_json(manifest_path, manifest)
    except OSError as exc:
        return {**result, 'error': f'Cannot record bundle outcome: {type(exc).__name__}'}
    return {**result, 'status': manifest['status'],
            'error': manifest.get('error') or manifest.get('visualization_error'),
            'human_held': manifest.get('human_held'),
            'artifacts': sorted(manifest['artifacts']),
            'routes': {name: route['state'] for name, route in manifest['routes'].items()}}


def prepare_and_deliver(cohort, mode, report_date, audience, rolling, programs, recipients,
                        options, directory, manifest):
    leadership = audience == 'leadership'
    healthcare = programs == ('Healthcare',)
    if (options.offline or options.dashboard_only) and not options.dry_run:
        raise ValueError('Offline and dashboard-only runs cannot send emails')
    profile = load_profile()
    profile_hash = hashlib.sha256(json.dumps(profile, sort_keys=True).encode()).hexdigest()
    if manifest.get('profile_sha256', profile_hash) != profile_hash:
        raise RuntimeError('Visualization profile changed for a saved run; use a new run root for a fresh preview')
    manifest['profile_sha256'] = profile_hash
    supported = bool(profile.get('enabled')) and str(cohort.id) in [str(x) for x in profile.get('cohorts', [])] and not healthcare
    needs_pdf = supported and mode in (0, 2)
    artifacts = manifest['artifacts']
    snapshot = directory / rolling.name
    if 'xlsx' in artifacts:
        verify_file(snapshot, artifacts['xlsx']['sha256'])
    else:
        if options.retry:
            raise RuntimeError('No completed report snapshot; retry must not requery SQL')
        source_hash = sha256(rolling)
        atomic_copy(rolling, snapshot)
        verify_file(snapshot, source_hash)
        if not options.offline:
            if leadership:
                from LeadershipReport.update_overall_report import main
                main(cohort.id, cohort.semester_byu, cohort.internship_years, str(snapshot), report_date=report_date)
            else:
                if not healthcare:
                    from ProgramReports.create_program_reports import main
                    main(programs, cohort, output_path=snapshot, report_date=report_date)
                from ProgramReports.placement_details import main
                main(programs, cohort, output_path=snapshot, report_date=report_date)
        if not healthcare:
            warnings = []
            read_history(snapshot, PAGES if leadership else (OVERALL,) + tuple(programs),
                         report_date, cohort.start_date, warnings)
            manifest['source_warnings'] = warnings
        else:
            from openpyxl import load_workbook
            wb = load_workbook(snapshot, read_only=True)
            try:
                if wb.sheetnames != ['Healthcare Full-Time', 'Healthcare Internships', 'Past Week Placement Details']:
                    raise ValueError('Healthcare sheet contract changed')
            finally:
                wb.close()
        artifacts['xlsx'] = {'file': snapshot.name, 'sha256': sha256(snapshot)}
        manifest['source_rolling_sha256'] = source_hash
        write_json(directory / 'manifest.json', manifest)
    # Promote valid data independently of PDF success; never overwrite newer history.
    if not options.dry_run and not options.test_recipient and not manifest.get('history_promoted'):
        current = sha256(rolling)
        if current not in (manifest['source_rolling_sha256'], artifacts['xlsx']['sha256']):
            raise RuntimeError('Rolling workbook changed since staging; refusing to overwrite newer history')
        backup = rolling.with_name(rolling.stem + f'.before-{report_date}.xlsx')
        if not backup.exists() and current == manifest['source_rolling_sha256']:
            atomic_copy(rolling, backup)
        atomic_copy(snapshot, rolling)
        manifest['history_promoted'] = True
        write_json(directory / 'manifest.json', manifest)
    pdf = None
    manifest.pop('visualization_error', None)
    if supported:
        try:
            dashboard, metadata = prepare_dashboard(snapshot, directory, manifest, cohort,
                report_date, programs, options, profile, leadership)
            if needs_pdf and not options.dashboard_only:
                if 'pdf' in artifacts:
                    pdf = directory / artifacts['pdf']['file']
                    verify_file(pdf, artifacts['pdf']['sha256'])
                else:
                    label = 'Leadership' if leadership else '-'.join(programs)
                    pdf = directory / f'{cohort.id}-FullTime-{label}-{report_date}.pdf'
                    info = render_pdf(dashboard, pdf, pages=metadata['pages'], year=cohort.id,
                        report_date=report_date, **{k: profile[k] for k in (
                            'cover', 'dpi', 'timeout', 'print_area', 'max_pdf_bytes') if k in profile})
                    artifacts['pdf'] = {'file': pdf.name, 'sha256': sha256(pdf), **info}
                    write_json(directory / 'manifest.json', manifest)
        except Exception as exc:
            pdf = None
            manifest['visualization_error'] = f'{type(exc).__name__}: {exc}'
            logger.error('%s/%s visualization failed: %s', cohort.id, audience, manifest['visualization_error'])
    settings = {}
    if not options.offline:
        from .database import load_environment
        settings = load_environment(ROOT / ('LeadershipReport' if leadership else 'ProgramReports'))
    prepare_routes(cohort, mode, report_date, audience, programs, recipients, options,
                   directory, manifest, snapshot, pdf, needs_pdf, profile, settings=settings)
    deliver(directory, manifest, dry_run=options.dry_run, resend_uncertain=options.resend_uncertain,
            allowed_channels=('box', 'zoho') if manifest.get('human_held') else None, settings=settings)


def prepare_dashboard(snapshot, directory, manifest, cohort, report_date, programs, options, profile, leadership):
    artifacts = manifest['artifacts']
    if 'dashboard' in artifacts:
        dashboard = directory / artifacts['dashboard']['file']
        verify_file(dashboard, artifacts['dashboard']['sha256'])
        verify_file(dashboard.with_suffix('.json'), artifacts['dashboard']['metadata_sha256'])
        return dashboard, json.loads(dashboard.with_suffix('.json').read_text())
    pointer_path = Path(options.run_root) / 'state' / namespace_for(options) / f'{cohort.id}.json'
    base = None
    if leadership and pointer_path.exists():
        pointer = json.loads(pointer_path.read_text())
        if pointer['report_date'] > str(report_date):
            raise RuntimeError('Newer dashboard state exists; retry the older saved dashboard instead')
        base = Path(options.run_root) / pointer['path']
        verify_file(base, pointer['sha256'])
        verify_file(base.with_suffix('.json'), pointer['metadata_sha256'])
    dashboard, _, metadata = build_visualization(snapshot, directory, cohort=cohort,
        report_date=report_date, programs=programs, base=base, profile=profile, render=False,
        source_root=options.source_root)
    artifacts['dashboard'] = {'file': dashboard.name, 'sha256': sha256(dashboard),
                              'metadata_sha256': sha256(dashboard.with_suffix('.json'))}
    write_json(directory / 'manifest.json', manifest)
    if leadership:
        write_json(pointer_path, {'report_date': str(report_date),
            'path': str(dashboard.relative_to(Path(options.run_root))), 'sha256': sha256(dashboard),
            'metadata_sha256': sha256(dashboard.with_suffix('.json'))})
    return dashboard, metadata


def prepare_routes(cohort, mode, report_date, audience, programs, recipients, options,
                   directory, manifest, snapshot, pdf, needs_pdf, profile, settings=None):
    leadership = audience == 'leadership'
    cc = env_list('CC_ADDRS', settings) if leadership else []
    bcc = env_list('BCC_ADDRS', settings) if leadership else []
    to = env_list('TO_ADDRS', settings) if leadership else recipients
    archive = (cohort.monthly_box_leadership if leadership else cohort.monthly_box_programs) if mode in (1, 2) else (
        cohort.weekly_box_leadership if leadership else cohort.weekly_box_programs)
    key = f'{cohort.id}/{report_date}/{mode}/{audience}'
    def save(channel, msg, envelope):
        save_route(directory, manifest, channel, msg, envelope,
                   max_bytes=profile.get('max_message_bytes', 25000000), test_recipient=options.test_recipient)
    save('box', message(f'Auto upload: {snapshot.stem} {report_date}', 'Automated report archive.',
         [snapshot], [archive], message_key=key + '/box'), [archive])
    if leadership:
        zoho = env_list('Zoho_Box', settings)
        save('zoho', message(f'Auto upload: {snapshot.stem} {report_date}', 'Automated workbook import.',
             [snapshot], zoho, message_key=key + '/zoho'), zoho)
    human_allowed = mode in cohort.human_modes
    manifest['human_policy_enabled'] = human_allowed
    manifest['pdf_required'] = needs_pdf
    if human_allowed and (not needs_pdf or pdf is not None):
        attachments = [snapshot] + ([pdf] if pdf is not None else [])
        if leadership:
            kind = 'Month End' if mode in (1, 2) else 'Weekly'
            subject = f'Class of {cohort.id} {kind} Placement Report - {report_date}'
            body = f'Good Morning,\n\nAttached is the Class of {cohort.id} placement report as of {report_date}. The workbook includes full-time and internship totals, program histories, and placement summaries.'
            if pdf:
                body += '\n\nThe attached PDF contains full-time dashboards generated from this report. Current-month percentages are snapshots as of the report date; prior-year comparisons retain recorded observations.'
            body += '\n\nPlease contact the BCC Data Team with any discrepancies.\n\nBCC Data Team'
        else:
            from ProgramReports.email_templates import subject_and_body
            subject, body = subject_and_body(audience, ', '.join(programs), mode, report_date,
                                            cohort.id, has_visualization=pdf is not None)
        save('human', message(subject, body, attachments, to, cc=cc, message_key=key + '/human'), list(to) + cc + bcc)
        manifest.pop('human_held', None)
    elif human_allowed:
        manifest['human_held'] = 'Dashboard-only preview' if options.dashboard_only else 'Required PDF unavailable'
    else:
        manifest['human_held'] = 'Disabled for this report mode'
    write_json(directory / 'manifest.json', manifest)
