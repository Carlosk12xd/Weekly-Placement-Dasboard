"""Weekly cross-checks of a finished report workbook, read back from the saved file.

Errors mean the workbook is wrong and must not be sent or promoted:
  - this week's column differs from the SQL counts fetched for it, or a status returned
    by SQL has no row in the table (those students would be missing from Class Size);
  - a Weekly History column is dated outside the class window, duplicated, out of order,
    or the newest column is not the report date;
  - Class Size is not the sum of the statuses, or % Placed is not
    accepted / (accepted + actively seeking + not reported), in ANY column;
  - a Most Recent Friday table differs from the newest Weekly History column;
  - (leadership) the Summary sheet's accepted / seeking / not reported / % placed differ
    from the program's own table.
Warnings need a person to look but do not block: missed Fridays, a class size that moved
more than 5% in a week, accepted offers that went down, non-Friday columns, unreadable
historical headers, students whose graduation term is outside August Y-1..June Y.
"""
import re
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from .academic_year import in_window, graduation_terms, later_terms
from .history_window import header_date

ACCEPTED, SEEKING, NOT_REPORTED = 'Accepted an offer', 'Actively seeking', 'Not Reported'
TOTAL_LABELS = ('class size', 'total')
RATE_LABELS = ('% placed', 'placement %')
CLASS_SIZE_JUMP = 0.05


def is_history(name):
    """Weekly History tables: *_wh, or an even trailing number (BSAcc2, MBA_2, Class4, BSFin6)."""
    if name.lower().endswith('wh'):
        return True
    digit = re.search(r'(\d)$', name)
    return bool(digit) and int(digit[1]) % 2 == 0


def history_name(mrf):
    if mrf.lower().endswith('mrf'):
        return mrf[:-3] + 'wh'
    return mrf[:-1] + str(int(mrf[-1]) + 1)


def placement_rate(acc, seek, nr):
    """Same arithmetic as the builders: percent rounded to 2 places, stored as a fraction."""
    return round(acc * 100.0 / (acc + seek + nr), 2) / 100 if acc + seek + nr else 0


def count(value):
    if value in (None, '', '-'):
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or int(value) != value:
        raise ValueError(repr(value))
    return int(value)


def read_table(ws, table):
    """{'name', 'sheet', 'rows': {label: row}, 'total', 'rate', 'columns': [(header, {label: value})]}"""
    left, top, right, bottom = range_boundaries(table.ref)
    rows, total, rate = {}, None, None
    for r in range(top + 1, bottom + 1):
        label = str(ws.cell(r, left).value or '').strip()
        if label.lower() in TOTAL_LABELS:
            total = r
        elif label.lower() in RATE_LABELS:
            rate = r
        elif label:
            rows[label] = r
    columns = []
    for c in range(left + 1, right + 1):
        values = {label: ws.cell(r, c).value for label, r in rows.items()}
        values['__total'] = ws.cell(total, c).value if total else None
        values['__rate'] = ws.cell(rate, c).value if rate else None
        columns.append((ws.cell(top, c).value, values))
    return {'name': table.displayName, 'sheet': ws.title, 'rows': rows, 'columns': columns,
            'has_totals': total is not None and rate is not None,
            'full_time': 'job search' in str(ws.cell(top, left).value or '').lower()}


def check_column(table, header, values, errors):
    where = f"{table['sheet']}!{table['name']} {header}"
    try:
        counts = {label: count(values[label]) for label in table['rows']}
    except ValueError as exc:
        errors.append(f'{where}: count is not a whole number: {exc}')
        return None
    if not table['has_totals']:
        return counts
    if values['__total'] != sum(counts.values()):
        errors.append(f"{where}: Class Size {values['__total']} != sum of statuses {sum(counts.values())}")
    acc, seek, nr = counts.get(ACCEPTED, 0), counts.get(SEEKING, 0), counts.get(NOT_REPORTED, 0)
    expected = placement_rate(acc, seek, nr)
    rate = values['__rate']
    if not isinstance(rate, (int, float)) or abs(rate - expected) > 0.00005001:
        errors.append(f'{where}: % Placed {rate!r} != {acc}/({acc}+{seek}+{nr}) = {expected:.4f}')
    return counts


def check_report(path, *, class_year, report_date, sql=None, pruned=None, population=None,
                 population_error=None):
    errors, warnings, info = [], [], []
    label = report_date.strftime('%m/%d/%Y')
    wb = load_workbook(path)
    try:
        tables = {}
        for ws in wb.worksheets:
            for name in ws.tables:
                table = ws.tables[name]
                left, top, _, _ = range_boundaries(table.ref)
                if 'search status' in str(ws.cell(top, left).value or '').lower():
                    tables[name] = read_table(ws, table)
        if not tables:
            errors.append('No placement tables found')
        newest = {}
        gaps_reported = False
        for name, table in tables.items():
            if not is_history(name):
                continue
            dated = []
            for header, values in table['columns']:
                when = header_date(header)
                if when is None:
                    warnings.append(f'{table["sheet"]}!{name}: column {header!r} is not a date; kept as is')
                    continue
                if not in_window(class_year, when):
                    errors.append(f'{table["sheet"]}!{name}: {header} is outside the Class of {class_year} window')
                counts = check_column(table, header, values, errors)
                dated.append((when, header, counts, values))
            whens = [d[0] for d in dated]
            if len(set(whens)) != len(whens):
                errors.append(f'{table["sheet"]}!{name}: the same date appears in more than one column')
            if whens != sorted(whens):
                errors.append(f'{table["sheet"]}!{name}: Weekly History columns are not in date order')
            if not dated or dated[-1][1] != label:
                errors.append(f'{table["sheet"]}!{name}: newest column is '
                              f'{dated[-1][1] if dated else "missing"}, expected {label}')
                continue
            newest[name] = dated[-1][2]
            if not gaps_reported:  # dates are shared by every table; report them once
                gaps_reported = True
                odd = sorted({d[1] for d in dated if d[0].weekday() != 4})
                if odd:
                    warnings.append(f'Non-Friday history columns (kept): {", ".join(odd)}')
                for (before, *_), (after, *_) in zip(dated, dated[1:]):
                    missed = (after - before).days // 7 - 1
                    if missed > 0:
                        warnings.append(f'No report for {missed} Friday(s) between {before:%m/%d/%Y} and {after:%m/%d/%Y}; '
                                        'the trend line jumps across that gap')
            if len(dated) > 1 and dated[-1][2] and dated[-2][2] and table['has_totals'] and table['full_time']:
                previous, current = dated[-2][3]['__total'], dated[-1][3]['__total']
                if isinstance(previous, int) and previous and abs(current - previous) / previous > CLASS_SIZE_JUMP:
                    warnings.append(f'{table["sheet"]}!{name}: Class Size moved {previous} -> {current} '
                                    f'since {dated[-2][1]}')
                if dated[-1][2].get(ACCEPTED, 0) < dated[-2][2].get(ACCEPTED, 0):
                    warnings.append(f'{table["sheet"]}!{name}: Accepted went down '
                                    f'{dated[-2][2].get(ACCEPTED, 0)} -> {dated[-1][2].get(ACCEPTED, 0)} since {dated[-2][1]}')
        recent = {}
        for name, table in tables.items():
            if is_history(name):
                continue
            header, values = table['columns'][0] if table['columns'] else (None, {})
            if header != label:
                errors.append(f'{table["sheet"]}!{name}: Most Recent Friday header is {header!r}, expected {label}')
            counts = check_column(table, header, values, errors) if table['columns'] else None
            recent[name] = counts
            partner = newest.get(history_name(name))
            if counts is not None and partner is not None and counts != partner:
                errors.append(f'{table["sheet"]}!{name} differs from {history_name(name)} for {label}')
        for name, expected in (sql or {}).items():
            table = tables.get(name)
            if table is None:
                errors.append(f'{name}: table missing from the workbook')
                continue
            actual = newest.get(name) if is_history(name) else recent.get(name)
            if actual is None:
                continue  # already reported above
            missing = {s: n for s, n in expected.items() if s not in table['rows'] and n}
            if missing:
                errors.append(f'{table["sheet"]}!{name}: SQL statuses with no row in the table '
                              f'(left out of Class Size): {missing}')
            wrong = {s: (actual[s], expected.get(s, 0)) for s in actual if actual[s] != expected.get(s, 0)}
            if wrong:
                errors.append(f'{table["sheet"]}!{name}: workbook != SQL for {label} (workbook, SQL): {wrong}')
        if 'Summary - Full Time' in wb.sheetnames:
            check_summary(wb['Summary - Full Time'], newest, errors, warnings)
    finally:
        wb.close()
    for name, removed in (pruned or {}).items():
        info.append(f'{name}: removed {removed if isinstance(removed, int) else len(removed)} '
                    f'entries outside the Class of {class_year} window or duplicated'
                    + ('' if isinstance(removed, int) else f' ({", ".join(map(str, removed[:6]))}{"..." if len(removed) > 6 else ""})'))
    if population_error:
        warnings.append(f'Population cross-check query failed: {population_error}')
    if population:
        check_population(population, class_year, errors, warnings, info)
    return {'errors': errors, 'warnings': warnings, 'info': info, 'tables_checked': len(tables)}


def check_summary(ws, newest, errors, warnings):
    from LeadershipReport.update_overall_report import byprog_full_names
    for row in ws.iter_rows(min_row=3, max_col=9, values_only=True):
        program = row[0]
        if not isinstance(program, str):
            continue
        try:
            history = byprog_full_names(program)[1]
        except Exception:
            continue
        counts = newest.get(history)
        if counts is None:
            continue
        acc, seek, nr = counts.get(ACCEPTED, 0), counts.get(SEEKING, 0), counts.get(NOT_REPORTED, 0)
        if (row[2], row[3], row[4]) != (acc, seek, nr):
            errors.append(f'Summary {program}: accepted/seeking/not reported {row[2:5]} != {history} {(acc, seek, nr)}')
        rate = placement_rate(acc, seek, nr)
        if not isinstance(row[1], (int, float)) or abs(row[1] - rate) > 0.00005001:
            errors.append(f'Summary {program}: % Placed {row[1]!r} != {rate:.4f} from {history}')
        if isinstance(row[8], int) and row[8] != sum(counts.values()):
            warnings.append(f'Summary {program}: Total {row[8]} (summary query) != Class Size {sum(counts.values())} '
                            '(program query)')


def check_population(rows, class_year, errors, warnings, info):
    """rows: (class_of, graduation term, enroll status, count) for the full-time population.
    The query already moves students graduating after June Y to a later class, so anyone left
    outside the window was kept by the class_of rule (still enrolled past their term)."""
    terms = graduation_terms(class_year)
    later = set(later_terms(class_year))
    outside, unknown, total = {}, 0, 0
    for class_of, term, status, n in rows:
        total += int(n)
        term = str(term or '').strip()
        if not re.fullmatch(r'(Winter|Spring|Summer|Fall) \d{4}', term):
            unknown += int(n)
        elif term in later:
            errors.append(f'{n} students graduating {term} (class_of {class_of}) are counted in Class of '
                          f'{class_year}; they belong to a later class')
        elif term not in terms:
            key = f'class_of {class_of}, graduating {term}, {status}'
            outside[key] = outside.get(key, 0) + int(n)
    info.append(f'Population: {total} students; graduation terms counted as Class of {class_year}: '
                f'{", ".join(sorted(terms))}')
    if outside:
        warnings.append(f'{sum(outside.values())} of {total} students in the report graduate '
                        f'outside August {int(class_year) - 1}-June {class_year}: '
                        + '; '.join(f'{k}: {v}' for k, v in sorted(outside.items())))
    if unknown:
        warnings.append(f'{unknown} students have no graduation term like "Fall 2026"; kept by the class_of rule')


def compare_reports(leadership, director, *, report_date, programs):
    """Same-day leadership vs director workbook. Queried minutes apart, so differences are warnings."""
    from Visualization.update_dashboard import read_history, OVERALL
    warnings = []
    pages = (OVERALL,) + tuple(programs)
    a = read_history(leadership, pages, report_date, report_date)
    b = read_history(director, pages, report_date, report_date)
    for page in pages:
        if a[page][report_date]['counts'] != b[page][report_date]['counts']:
            warnings.append(f'{page}: director workbook counts differ from the leadership report for '
                            f'{report_date:%m/%d/%Y} ({b[page][report_date]["rate"]:.2%} vs {a[page][report_date]["rate"]:.2%})')
    return warnings
