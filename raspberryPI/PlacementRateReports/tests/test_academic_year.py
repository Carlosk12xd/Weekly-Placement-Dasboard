"""Class window, October rollover, new-class workbooks, weekly cross-checks and the
Friday visualization step. SQL is mocked (tests/october_fixture.py); SMTP is never called."""
from datetime import date
import json
import sqlite3
from pathlib import Path
import shutil
import tempfile
from unittest import TestCase, main
from unittest.mock import patch
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from core.academic_year import (window_start, window_end, report_class_year, graduation_terms, last_friday,
                                full_time_population_sql)
from core.cohort_manager import CohortManager, derive_cohort
from core.daily_runner import load_cohorts
from core.history_window import header_date
from core.pipeline import execute_bundle
from core.report_artifacts import ROOT, RunOptions
from core.report_checks import check_report, check_population
from core.rollover import create_class_workbook, class_report_path
from Visualization.update_dashboard import read_history, OVERALL
from tests.october_fixture import build_week

FIXTURES = ROOT / 'tests/fixtures'
LEADERSHIP = 'LeadershipReport/Reports/2027_weekly_placement_report.xlsx'
FINANCE = 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
SYSTEMS = 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSIS-MISM.xlsx'
OCT9 = date(2026, 10, 9)


def history_headers(path, sheet, table):
    wb = load_workbook(path)
    try:
        ws = wb[sheet]
        left, top, right, _ = range_boundaries(ws.tables[table].ref)
        return [ws.cell(top, c).value for c in range(left + 1, right + 1)]
    finally:
        wb.close()


class CalendarTests(TestCase):
    def setUp(self):
        self.cohorts = load_cohorts()[1]
        self.manager = CohortManager(self.cohorts)

    def test_window_and_report_class(self):
        self.assertEqual(last_friday(2026, 6), date(2026, 6, 26))
        self.assertEqual(window_start(2027), date(2026, 6, 22))   # matches cohorts.yml start_date
        self.assertEqual(window_start(2026), date(2025, 6, 23))
        self.assertEqual(window_end(2027), date(2027, 9, 30))
        self.assertEqual(last_friday(2027, 9), date(2027, 9, 24))  # last column of the Class of 2027
        self.assertEqual(report_class_year(date(2026, 9, 30)), 2026)
        self.assertEqual(report_class_year(date(2026, 10, 1)), 2027)
        self.assertEqual(report_class_year(date(2027, 6, 25)), 2027)
        self.assertEqual(graduation_terms(2027), {'Summer 2026', 'Fall 2026', 'Winter 2027', 'Spring 2027'})

    def test_full_time_population_rule(self):
        students = [  # name, class_of, enroll_status, expected term, actual term
            ('class 2027 enrolled', 2027, 'Enrolled', 'Winter 2027', None),
            ('class 2027 graduated', 2027, 'Graduated', 'Fall 2026', 'Fall 2026'),
            ('class 2027, no term', 2027, 'Enrolled', None, None),
            ('class 2026 still enrolled', 2026, 'Enrolled', 'Spring 2026', None),
            ('class 2025 still enrolled', 2025, 'Enrolled', 'Fall 2026', ''),
            ('class 2024 still enrolled', 2024, 'Enrolled', 'Winter 2027', None),
            ('class 2026 graduated', 2026, 'Graduated', 'Spring 2026', 'Spring 2026'),
            ('class 2027, graduating Aug 2027', 2027, 'Enrolled', 'Summer 2027', None),
            ('class 2026 enrolled, graduating Dec 2027', 2026, 'Enrolled', 'Fall 2027', None),
            ('class 2025, graduated Aug 2026', 2025, 'Graduated', 'Spring 2026', 'Summer 2026'),
            ('class 2028, graduating Apr 2027', 2028, 'Enrolled', 'Winter 2027', None),
            ('class 2028', 2028, 'Enrolled', 'Winter 2028', None),
            ('class 2028, graduated Apr 2027', 2028, 'Graduated', 'Winter 2028', 'Winter 2027'),
        ]
        db = sqlite3.connect(':memory:')
        db.execute('CREATE TABLE s (name, class_of, enroll_status, exp_grad_semester, act_grad_semester)')
        db.executemany('INSERT INTO s VALUES (?, ?, ?, ?, ?)', students)
        counted = lambda year: {r[0] for r in db.execute(f'SELECT name FROM s WHERE {full_time_population_sql(year)}')}
        self.assertEqual(counted(2027), {
            'class 2027 enrolled', 'class 2027 graduated', 'class 2027, no term', 'class 2026 still enrolled',
            'class 2025 still enrolled', 'class 2024 still enrolled', 'class 2025, graduated Aug 2026',
            'class 2028, graduating Apr 2027', 'class 2028, graduated Apr 2027'})
        # 2028 keeps students still enrolled past their term (class 2027 enrolled, class 2026 still
        # enrolled), but not anyone who already graduated before August 2027.
        self.assertEqual(counted(2028), {
            'class 2027, graduating Aug 2027', 'class 2026 enrolled, graduating Dec 2027', 'class 2028',
            'class 2027 enrolled', 'class 2027, no term', 'class 2026 still enrolled', 'class 2028, graduating Apr 2027'})
        self.assertEqual(counted(2026), {'class 2026 still enrolled', 'class 2026 graduated'})
        self.assertIn("a.class_of = 2027", full_time_population_sql(2027, 'a.'))

    def test_active_classes_overlap_from_late_june_through_september(self):
        ids = lambda day: [c.id for c in self.manager.get_active_cohorts(day)]
        self.assertEqual(ids(date(2026, 6, 19)), ['2026'])
        self.assertEqual(ids(date(2026, 6, 26)), ['2026', '2027'])
        self.assertEqual(ids(date(2026, 9, 25)), ['2026', '2027'])   # last Friday of September
        self.assertEqual(ids(date(2026, 10, 2)), ['2027'])
        self.assertEqual(ids(date(2027, 6, 25)), ['2027', '2028'])   # 2028 derived automatically
        self.assertEqual(ids(date(2027, 10, 1)), ['2028'])

    def test_next_class_is_derived_from_the_latest_configured(self):
        latest = next(c for c in self.cohorts if c.id == '2027')
        nxt = derive_cohort(latest, 2028)
        self.assertEqual((nxt.id, nxt.label), ('2028', 'Class of 2028'))
        self.assertEqual((nxt.start_date, nxt.grad_date), (date(2027, 6, 21), date(2028, 6, 26)))
        self.assertEqual(nxt.internship_years, ['2029', '2030', '2031'])
        self.assertEqual(nxt.semester_byu, ['20285', '20295', '20305'])
        self.assertEqual(nxt.overview_template, 'LeadershipReport/Reports/2028_weekly_placement_report.xlsx')
        self.assertEqual(nxt.weekly_box_leadership, latest.weekly_box_leadership)
        self.assertEqual(nxt.human_modes, latest.human_modes)
        self.assertEqual(class_report_path(Path('/pi/2027_Placement_Reports/2027-WeeklyPlacement-MBA.xlsx'), 2027, 2028),
                         Path('/pi/2028_Placement_Reports/2028-WeeklyPlacement-MBA.xlsx'))


class WeeklyHistoryTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.temp.name)
        cls.cohort = next(c for c in load_cohorts()[1] if c.id == '2027')
        cls.lead_capture, cls.fin_capture = {}, {}
        cls.leadership = build_week(FIXTURES / LEADERSHIP, cls.folder / LEADERSHIP, cohort=cls.cohort,
                                    report_date=OCT9, capture=cls.lead_capture)
        cls.finance = build_week(FIXTURES / FINANCE, cls.folder / FINANCE, cohort=cls.cohort,
                                 report_date=OCT9, programs=('BSFin',), capture=cls.fin_capture)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_history_keeps_only_the_window_one_column_per_date(self):
        for path, sheet, table in [(self.leadership, 'Total - Full Time', 'FT_total_wh'),
                                   (self.leadership, 'By Program - Internships', 'MBA_int2'),
                                   (self.finance, '2027 MSB Overall', 'Class2'),
                                   (self.finance, 'BSFin', 'BSFin6')]:
            headers = history_headers(path, sheet, table)
            dates = [header_date(h) for h in headers]
            self.assertTrue(all(window_start(2027) <= d <= window_end(2027) for d in dates), headers)
            self.assertEqual(len(set(dates)), len(dates))
            self.assertEqual(dates, sorted(dates))
            self.assertEqual(headers[-1], '10/09/2026')
            self.assertNotIn('07/10/2026_2', headers)
        self.assertEqual(history_headers(self.leadership, 'Total - Full Time', 'FT_total_wh')[0], '06/26/2026')
        wb = load_workbook(self.leadership, read_only=True)
        try:
            zoho_dates = {row[3].date() for row in wb['ZOHO_DATA'].iter_rows(min_row=2, values_only=True) if row[0]}
        finally:
            wb.close()
        self.assertEqual(min(zoho_dates), date(2026, 6, 26))
        self.assertIn(OCT9, zoho_dates)

    def test_new_week_rate_matches_counts_and_checks_pass(self):
        records = read_history(self.leadership, (OVERALL, 'MBA'), OCT9, self.cohort.start_date)
        counts = records[OVERALL][OCT9]['counts']
        self.assertEqual(records[OVERALL][OCT9]['rate'],
                         round(counts[0] * 100 / (counts[0] + counts[1] + counts[3]), 2) / 100)
        for path, capture in [(self.leadership, self.lead_capture), (self.finance, self.fin_capture)]:
            result = check_report(path, class_year=2027, report_date=OCT9, sql=capture['sql'],
                                  pruned=capture['pruned'], population=capture.get('population'))
            self.assertEqual(result['errors'], [], path.name)
            self.assertTrue(any('6 Friday(s)' in w for w in result['warnings']))
        self.assertEqual(len(self.lead_capture['sql']), 4 + 4 * 14)

    def tampered(self, cell, value, sheet='By Program - Full Time'):
        copy = self.folder / f'tampered-{cell}.xlsx'
        wb = load_workbook(self.leadership)
        wb[sheet][cell] = value
        wb.save(copy)
        return check_report(copy, class_year=2027, report_date=OCT9, sql=self.lead_capture['sql'])['errors']

    def test_checks_catch_wrong_counts_rates_and_dates(self):
        wb = load_workbook(self.leadership)
        ws = wb['By Program - Full Time']
        left, top, right, bottom = range_boundaries(ws.tables['BSAcc2'].ref)
        newest, older = ws.cell(top + 1, right).coordinate, ws.cell(bottom, left + 2).coordinate
        header = ws.cell(top, left + 2).coordinate
        wb.close()
        errors = self.tampered(newest, 999)  # accepted this week
        self.assertTrue(any('workbook != SQL' in e for e in errors), errors)
        self.assertTrue(any('Class Size' in e for e in errors), errors)
        self.assertTrue(any('Summary BSAcc' in e for e in errors), errors)
        self.assertTrue(any('% Placed' in e for e in self.tampered(older, 0.99)))  # an earlier week's rate
        self.assertTrue(any('outside the Class of 2027 window' in e for e in self.tampered(header, '05/01/2026')))
        extra = {**self.lead_capture['sql'], 'FT_total_wh': {**self.lead_capture['sql']['FT_total_wh'], 'New status': 4}}
        errors = check_report(self.leadership, class_year=2027, report_date=OCT9, sql=extra)['errors']
        self.assertTrue(any('left out of Class Size' in e for e in errors), errors)

    def test_population_outside_graduation_window_is_reported(self):
        errors, warnings, info = [], [], []
        check_population([(2027, 'Winter 2027', 'Enrolled', 10), (2025, 'Spring 2025', 'Enrolled', 2),
                          (2027, None, 'Enrolled', 1)], 2027, errors, warnings, info)
        self.assertEqual(errors, [])
        self.assertIn('2 of 13 students', warnings[0])
        self.assertIn('class_of 2025, graduating Spring 2025', warnings[0])
        self.assertIn('1 students have no graduation term', warnings[1])
        check_population([(2027, 'Summer 2027', 'Enrolled', 4)], 2027, errors, [], [])
        self.assertIn('4 students graduating Summer 2027', errors[0])


class RolloverTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.cohort28 = CohortManager(load_cohorts()[1]).cohort(2028)
        self.day = date(2027, 6, 25)  # first Friday of the Class of 2028

    def tearDown(self):
        self.temp.cleanup()

    def test_new_program_workbook_names_the_next_class_and_starts_empty(self):
        new = create_class_workbook(FIXTURES / FINANCE, self.folder / 'new.xlsx', old_year=2027, new_year=2028)
        wb = load_workbook(new)
        self.assertEqual(wb.sheetnames, ['2028 MSB Overall', 'BSFin', 'BSFin Placement Details'])
        self.assertEqual(wb['2028 MSB Overall']['A1'].value, '2028 MSB Overall Placement')
        self.assertEqual(wb['BSFin']['A33'].value, 'CLASS OF 2029 INTERNSHIP PLACEMENT')
        self.assertEqual(wb['BSFin Placement Details']['N1'].value, 'BSFin 2029')
        self.assertEqual(wb['BSFin'].tables['BSFin2'].ref, 'A18:A30')
        wb.close()
        capture = {}
        built = build_week(new, self.folder / '2028-WeeklyPlacement-BSFin.xlsx', cohort=self.cohort28,
                           report_date=self.day, programs=('BSFin',), capture=capture, counts_from=FIXTURES / FINANCE)
        self.assertEqual(history_headers(built, 'BSFin', 'BSFin2'), ['06/25/2027'])
        self.assertEqual(check_report(built, class_year=2028, report_date=self.day, sql=capture['sql'])['errors'], [])
        self.assertIn(self.day, read_history(built, (OVERALL, 'BSFin'), self.day, self.cohort28.start_date)['BSFin'])

    def test_new_leadership_workbook_has_only_its_own_weeks(self):
        new = create_class_workbook(FIXTURES / LEADERSHIP, self.folder / 'lead.xlsx', old_year=2027, new_year=2028)
        capture = {}
        built = build_week(new, self.folder / '2028_weekly_placement_report.xlsx', cohort=self.cohort28,
                           report_date=self.day, capture=capture, counts_from=FIXTURES / LEADERSHIP)
        self.assertEqual(history_headers(built, 'Total - Full Time', 'FT_total_wh'), ['06/25/2027'])
        self.assertEqual(check_report(built, class_year=2028, report_date=self.day, sql=capture['sql'])['errors'], [])
        wb = load_workbook(built, read_only=True)
        dates = {row[3].date() for row in wb['ZOHO_DATA'].iter_rows(min_row=2, values_only=True) if row[0]}
        wb.close()
        self.assertEqual(dates, {self.day})

    def test_pipeline_creates_missing_class_workbook_without_touching_production(self):
        source = self.folder / 'source'
        target = source / 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
        target.parent.mkdir(parents=True)
        shutil.copy2(FIXTURES / FINANCE, target)
        rolling = class_report_path(target, 2027, 2028)
        from ProgramReports import create_program_reports
        real_main = create_program_reports.main
        def builder(programs, cohort, output_path=None, report_date=None, capture=None):
            build_week(output_path, output_path, cohort=cohort, report_date=report_date, programs=programs,
                       capture=capture, counts_from=FIXTURES / FINANCE, builder_main=real_main)
        with patch('ProgramReports.create_program_reports.main', side_effect=builder), \
             patch('ProgramReports.placement_details.main'), patch('core.pipeline.date') as today, \
             patch('smtplib.SMTP', side_effect=AssertionError('no SMTP')):
            today.today.return_value = self.day
            result = execute_bundle(self.cohort28, 0, self.day, 'Tanya', rolling, ('BSFin',), ['x@example.invalid'],
                                    RunOptions(dry_run=True, run_root=self.folder / 'runs', dashboard_only=True))
        self.assertNotEqual(result['status'], 'failed', result)
        manifest = json.loads((Path(result['directory']) / 'manifest.json').read_text())
        self.assertEqual(manifest['created_from']['file'], '2027-WeeklyPlacement-BSFin.xlsx')
        self.assertEqual(manifest['checks']['errors'], 0)
        self.assertIn('archived only', manifest['human_held'])  # Class of 2028 is not the report class in June 2027
        self.assertFalse(rolling.exists())  # a dry run never writes the rolling workbook


class FridayStepTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.temp.name)
        cls.source = cls.folder / 'source'
        cls.cohort = next(c for c in load_cohorts()[1] if c.id == '2027')
        shutil.copytree(FIXTURES, cls.source)
        build_week(cls.source / LEADERSHIP, cls.source / LEADERSHIP, cohort=cls.cohort, report_date=OCT9)
        build_week(cls.source / FINANCE, cls.source / FINANCE, cohort=cls.cohort, report_date=OCT9, programs=('BSFin',))
        build_week(cls.source / SYSTEMS, cls.source / SYSTEMS, cohort=cls.cohort, report_date=OCT9,
                   programs=('BSIS', 'MISM'), shift=1)
        cls.options = RunOptions(offline=True, run_root=cls.folder / 'runs', source_root=cls.source, dashboard_only=True)
        cls.results = []
        with patch('smtplib.SMTP', side_effect=AssertionError('no SMTP')):
            for audience, rel, programs in [('leadership', LEADERSHIP, None), ('Tanya', FINANCE, ('BSFin',)),
                                            ('Bob', SYSTEMS, ('BSIS', 'MISM'))]:
                cls.results.append(execute_bundle(cls.cohort, 0, OCT9, audience, cls.source / rel, programs,
                                                  ['x@example.invalid'], cls.options))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    @staticmethod
    def fake_images(_, pages, folder, **kwargs):
        from PIL import Image
        Path(folder).mkdir(parents=True, exist_ok=True)
        shots = [Path(folder) / f'{i:02d}-{p}.png' for i, p in enumerate(pages, 1)]
        for shot in shots:
            Image.new('RGB', (400, 225), 'white').save(shot)
        return shots

    def test_october_bundles_send_to_people_and_pass_checks(self):
        for result in self.results:
            manifest = json.loads((Path(result['directory']) / 'manifest.json').read_text())
            self.assertEqual(manifest['checks']['errors'], 0, result)
            self.assertNotIn('archived only', manifest.get('human_held') or '')

    def test_friday_step_builds_checked_decks_from_this_runs_snapshots(self):
        from core import friday_visualization
        from visualization_test import run as runner
        with patch.object(runner, 'render_sheet_images', side_effect=self.fake_images), \
             patch('smtplib.SMTP', side_effect=AssertionError('no SMTP')):
            outcome = friday_visualization.run(OCT9, load_cohorts()[1], self.results, self.options)
        self.assertEqual(outcome['status'], 'preview', outcome)
        self.assertEqual(outcome['results'], {'leadership': 'preview', 'Tanya': 'preview', 'Bob': 'preview'})
        self.assertTrue(any(w.startswith('Bob: ') for w in outcome['comparison_warnings']))  # shift=1 vs 3
        self.assertFalse(any(w.startswith('Tanya: ') for w in outcome['comparison_warnings']))
        base = Path(outcome['run_root'])
        for audience, decks in [('leadership', 16), ('Tanya', 3), ('Bob', 4)]:
            manifest = json.loads((base / audience / 'manifest.json').read_text())
            self.assertEqual(manifest['class_years'], [2027])
            self.assertEqual(json.loads((base / audience / 'checks.json').read_text())['errors'], [])
            from pptx import Presentation
            self.assertEqual(len(Presentation(base / audience / manifest['deck']).slides), decks)
        self.assertTrue((base / 'source/LeadershipReport/Reports/2026_weekly_placement_report.xlsx').is_file())

    def test_dashboard_check_rejects_a_chart_that_disagrees_with_the_report(self):
        from visualization_test.checks import check_dashboard
        from visualization_test.dashboard import build_dashboard
        from Visualization.update_dashboard import Package, NITTY, PAGES
        out = self.folder / 'dash.xlsx'
        meta = build_dashboard(self.source / FINANCE, out, class_year=2027, as_of=OCT9,
                               start_date=self.cohort.start_date, programs=('BSFin',),
                               template=ROOT / 'visualization_test/assets/placement_dashboard_template.xlsx',
                               state_path=self.folder / 'state.json')
        kwargs = dict(pages=meta['pages'], class_year=2027, as_of=OCT9, start_date=self.cohort.start_date)
        self.assertEqual(check_dashboard(out, self.source / FINANCE, **kwargs)['errors'], [])
        package = Package(out)
        rate_row = 24 + 17 * PAGES.index('BSFin')
        last = len(meta['weekly_dates']) + 1
        from openpyxl.utils import get_column_letter
        package.set(NITTY, f'{get_column_letter(last)}{rate_row}', 0.5)
        package.set('BSFin', 'N24', '2026')
        package.save(self.folder / 'dash-bad.xlsx')
        errors = check_dashboard(self.folder / 'dash-bad.xlsx', self.source / FINANCE, **kwargs)['errors']
        self.assertTrue(any('10/09/2026: chart 0.5' in e for e in errors), errors)
        self.assertTrue(any('current class is' in e for e in errors), errors)


if __name__ == '__main__':
    main()
