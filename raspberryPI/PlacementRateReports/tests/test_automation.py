"""Offline regression tests; every database and SMTP execution is mocked."""
from datetime import date
from email import policy
from email.parser import BytesParser
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase, main
from unittest.mock import patch
import json
import os
import shutil
import tempfile
from contextlib import redirect_stdout
from io import StringIO
from zipfile import ZipFile

from core.daily_runner import load_cohorts, report_mode
from core.cohort_manager import CohortManager
from core.report_artifacts import ROOT, RunOptions, sha256
from core.delivery import message, save_route, deliver
from core.pipeline import execute_bundle
from Visualization.update_dashboard import (update_dashboard, Package, PAGES, NITTY,
    OVERALL, STATUSES, TEMPLATE, ContractError, monthly_row, read_history)
from Visualization.xlsx_package import NS

FIXTURES = ROOT / 'tests/fixtures'
REPORTS = FIXTURES if FIXTURES.exists() else ROOT


class DashboardTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.work = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.work.name)
        cls.source26 = REPORTS / 'LeadershipReport/Reports/2026_weekly_placement_report.xlsx'
        cls.source27 = REPORTS / 'LeadershipReport/Reports/2027_weekly_placement_report.xlsx'
        cls.june = cls.folder / 'june.xlsx'
        cls.new = cls.folder / '2027.xlsx'
        cls.m26 = update_dashboard(cls.source26, cls.june, cohort_year=2026,
            as_of=date(2026, 6, 26), start_date=date(2025, 6, 23))
        cls.m27 = update_dashboard(cls.source27, cls.new, cohort_year=2027,
            as_of=date(2026, 8, 21), start_date=date(2026, 6, 22))

    @classmethod
    def tearDownClass(cls):
        cls.work.cleanup()

    def test_june_counts_rates_and_monthly_cells(self):
        original, actual = Package(TEMPLATE), Package(self.june)
        self.assertEqual(len(self.m26['weekly_dates']), 37)
        for i, page in enumerate(PAGES):
            rate, header = 24 + 17 * i, 26 + 17 * i
            self.assertEqual(actual.get(NITTY, f'AL{rate}'), original.get(NITTY, f'AL{rate}'))
            for row in range(header + 1, header + 12):
                expected = original.get(NITTY, f'AL{row}')
                self.assertEqual(actual.get(NITTY, f'AL{row}'), 0 if expected == '-' else expected)
            self.assertEqual(actual.get(page, 'N35'), original.get(page, 'N35'))
        self.assertEqual(actual.get(OVERALL, 'N35'), 0.8581)

    def test_rollover_and_monthly_cycle(self):
        p = Package(self.new)
        old = Package(TEMPLATE)
        self.assertEqual(p.get(OVERALL, 'N24'), 2027)
        self.assertEqual(int(p.get(OVERALL, 'O24')), 2026)
        self.assertEqual(p.get(OVERALL, 'O35'), old.get(OVERALL, 'N35'))
        self.assertEqual(p.get(OVERALL, 'N25'), 0.1069)
        self.assertIsNone(p.get(OVERALL, 'N35'))
        self.assertEqual(monthly_row(2027, date(2026, 9, 25)), 26)
        self.assertEqual(monthly_row(2027, date(2027, 9, 17)), 38)
        self.assertIsNone(monthly_row(2027, date(2026, 6, 26)))
        self.assertIsNone(monthly_row(2027, date(2026, 7, 31)))

    def test_append_rerun_and_closed_month_finalization(self):
        next_path = self.folder / 'july.xlsx'
        meta = update_dashboard(self.source26, next_path, cohort_year=2026,
            as_of=date(2026, 7, 10), start_date=date(2025, 6, 23), base=self.june)
        again = self.folder / 'july_again.xlsx'
        rerun = update_dashboard(self.source26, again, cohort_year=2026,
            as_of=date(2026, 7, 10), start_date=date(2025, 6, 23), base=next_path)
        self.assertEqual(len(meta['weekly_dates']), 38)
        self.assertEqual(meta['weekly_dates'], rerun['weekly_dates'])
        self.assertEqual(meta['monthly'][OVERALL]['35']['date'], '2026-06-30')
        self.assertEqual(meta['monthly'][OVERALL]['35']['kind'], 'month_end')
        p = Package(again)
        self.assertEqual(p.get(NITTY, 'AM23'), '07/10/2026')
        for part in p.chart_parts():
            for formula in p.xml(part).findall('.//c:f', NS):
                if formula.text.startswith(NITTY + '!') and ':' in formula.text:
                    self.assertIn(':$AM$', formula.text)
                    for skip in p.xml(part).findall('.//c:catAx/c:tickLblSkip', NS):
                        self.assertEqual(int(skip.get('val')), 5)

    def test_chart_style_parts_and_cache_lengths(self):
        p = Package(self.new)
        self.assertEqual(len(p.chart_parts()), 45)
        with ZipFile(TEMPLATE) as original, ZipFile(self.new) as generated:
            for name in original.namelist():
                if name.startswith(('xl/charts/style', 'xl/charts/colors', 'xl/theme/')):
                    self.assertEqual(original.read(name), generated.read(name), name)
        for part in p.chart_parts():
            root = p.xml(part)
            for ref in root.findall('.//c:numRef', NS) + root.findall('.//c:strRef', NS):
                node = ref.find('c:f', NS)
                if node is None:
                    continue  # Literal/cache-only series labels are not range references.
                formula = node.text
                if formula.startswith(NITTY + '!') and ':' in formula:
                    count = ref.find('.//c:ptCount', NS)
                    self.assertEqual(int(count.get('val')), len(self.m27['weekly_dates']))

    def test_identical_duplicates_collapsed_and_malformed_header_reported(self):
        self.assertEqual(self.m27['weekly_dates'].count('2026-07-10'), 1)
        self.assertTrue(any(x['header'] == '10/31/20252' for x in self.m26['source_warnings']))

    def test_prior_year_refresh_uses_available_observations_without_lookahead(self):
        out = self.folder / 'historical_refresh.xlsx'
        meta = update_dashboard(self.source27, out, cohort_year=2027,
            as_of=date(2026, 8, 21), start_date=date(2026, 6, 22),
            historical_sources={'2026': self.source26})
        records = meta['historical_sources']['2026']['monthly'][OVERALL]
        self.assertEqual(records['35']['date'], '2026-06-30')
        self.assertEqual(records['36']['date'], '2026-07-31')
        self.assertEqual(records['37']['date'], '2026-08-21')
        self.assertEqual(records['37']['kind'], 'last_available')
        self.assertNotIn('38', records)
        self.assertIsNone(Package(out).get(OVERALL, 'O38'))

    def test_conflicting_duplicate_and_missing_current_date_fail(self):
        altered = self.folder / 'conflict.xlsx'
        p = Package(self.source27)
        p.set('Total - Full Time', 'Q18', 999)
        p.save(altered)
        with self.assertRaisesRegex(ContractError, 'conflicting duplicate'):
            read_history(altered, (OVERALL,), date(2026, 8, 21), date(2026, 6, 22))
        with self.assertRaisesRegex(ContractError, 'no observation'):
            read_history(self.source27, (OVERALL,), date(2026, 8, 28), date(2026, 6, 22))


class FlowTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.cohort = next(c for c in load_cohorts()[1] if c.id == '2027')

    def tearDown(self):
        self.temp.cleanup()

    def test_calendar_and_tracking_boundaries(self):
        self.assertEqual(report_mode(date(2026, 9, 25)), 0)
        self.assertEqual(report_mode(date(2026, 9, 30)), 1)
        self.assertEqual(report_mode(date(2027, 4, 30)), 2)
        self.assertIsNone(report_mode(date(2026, 9, 24)))
        manager = CohortManager([self.cohort])
        self.assertEqual(manager.get_active_cohorts(date(2027, 9, 19)), [self.cohort])
        self.assertEqual(manager.get_active_cohorts(date(2027, 9, 20)), [])
        self.assertEqual(self.cohort.human_modes, [0, 2])

    def test_finance_details_use_both_internship_years(self):
        from ProgramReports import placement_details
        original = REPORTS / 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
        output = self.folder / original.name
        shutil.copy2(original, output)
        class Cursor:
            def __init__(self): self.call = 0
            def execute(self, *args): self.call += 1
            def fetchall(self): return [[['FT', 'INTERN_2028', 'INTERN_2029'][self.call - 1]]]
            def close(self): pass
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        with patch.object(placement_details, 'connect_database', return_value=Connection()), \
             patch.object(placement_details, 'load_environment', return_value={'OUTPUT_PATH': str(output)}):
            placement_details.main(('BSFin',), self.cohort)
        p = Package(output)
        self.assertEqual(p.get('BSFin Placement Details', 'O5'), 'INTERN_2028')
        self.assertEqual(p.get('BSFin Placement Details', 'Z5'), 'INTERN_2029')

    def test_healthcare_body_is_not_overwritten(self):
        from ProgramReports.email_templates import subject_and_body
        for flag in (0, 1, 2):
            _, body = subject_and_body('Erica', 'Healthcare', flag, date(2026, 8, 21), '2027')
            self.assertIn('healthcare', body.lower())
            self.assertNotIn('Most Recent Friday', body)
            self.assertNotIn('next business day', body)

    def test_leadership_builder_explicit_date_path_and_same_day_reuse(self):
        from LeadershipReport import update_overall_report as builder
        from openpyxl import load_workbook
        source = REPORTS / 'LeadershipReport/Reports/2027_weekly_placement_report.xlsx'
        output = self.folder / source.name
        shutil.copy2(source, output)
        counts = [(s, n) for s, n in zip(STATUSES[:-1], [12, 8, 3, 2, 1, 1, 1, 1, 1, 1])]
        summary = [(p, 12, 8, 5, 6, 2, 31) for p in PAGES[1:]]
        cursor = SimpleNamespace(close=lambda: None)
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        with patch.object(builder, 'connect_database', return_value=conn), \
             patch.object(builder, 'fetch_rows', side_effect=lambda cur, sql: summary if 'AS offer_accepted' in sql else counts), \
             redirect_stdout(StringIO()):
            builder.main('2027', self.cohort.semester_byu, self.cohort.internship_years, str(output), report_date=date(2026, 9, 25))
            first = load_workbook(output)
            extent = first['Total - Full Time'].tables['FT_total_wh'].ref
            first.close()
            builder.main('2027', self.cohort.semester_byu, self.cohort.internship_years, str(output), report_date=date(2026, 9, 25))
        records = read_history(output, PAGES, date(2026, 9, 25), self.cohort.start_date)
        self.assertEqual(records[OVERALL][date(2026, 9, 25)]['rate'], 0.5455)
        second = load_workbook(output)
        self.assertEqual(second['Total - Full Time'].tables['FT_total_wh'].ref, extent)
        second.close()

    def test_program_builder_explicit_date_path_and_same_day_reuse(self):
        from ProgramReports import create_program_reports as builder
        from openpyxl import load_workbook
        source = REPORTS / 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
        output = self.folder / source.name
        shutil.copy2(source, output)
        counts = [(s, n) for s, n in zip(STATUSES[:-1], [12, 8, 3, 2, 1, 1, 1, 1, 1, 1])]
        cursor = SimpleNamespace(close=lambda: None)
        conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
        with patch.object(builder, 'connect_database', return_value=conn), \
             patch.object(builder, 'fetch_rows', return_value=counts), \
             patch.object(builder, 'load_environment', return_value={'OUTPUT_PATH': str(output)}), redirect_stdout(StringIO()):
            builder.main(('BSFin',), self.cohort, report_date=date(2026, 9, 25))
            first = load_workbook(output)
            extent = first['BSFin'].tables['BSFin2'].ref
            first.close()
            builder.main(('BSFin',), self.cohort, output_path=output, report_date=date(2026, 9, 25))
        records = read_history(output, (OVERALL, 'BSFin'), date(2026, 9, 25), self.cohort.start_date)
        self.assertEqual(records['BSFin'][date(2026, 9, 25)]['rate'], 0.5455)
        second = load_workbook(output)
        self.assertEqual(second['BSFin'].tables['BSFin2'].ref, extent)
        second.close()

    def test_pdf_failure_holds_human_and_retry_keeps_snapshot(self):
        from core import pipeline
        source = REPORTS / 'ProgramReports/2027_Placement_Reports/2027-WeeklyPlacement-BSFin.xlsx'
        options = RunOptions(offline=True, run_root=self.folder)
        with patch.object(pipeline, 'render_pdf', side_effect=RuntimeError('fixture conversion failure')), \
             patch('smtplib.SMTP', side_effect=AssertionError('SMTP must not be called')):
            result = execute_bundle(self.cohort, 0, date(2026, 8, 21), 'Tanya', source,
                                    ('BSFin',), ['recipient@example.invalid'], options)
        self.assertEqual(result['status'], 'failed')
        run = Path(result['directory'])
        manifest = json.loads((run / 'manifest.json').read_text())
        self.assertEqual(set(manifest['routes']), {'box'})
        self.assertIn('Required PDF', manifest['human_held'])
        before = manifest['artifacts']['xlsx']['sha256']
        def fake_render(_, output, **kwargs):
            from pypdf import PdfWriter
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with Path(output).open('wb') as stream: writer.write(stream)
            return {'pages': 1}
        options.retry = True
        with patch.object(pipeline, 'render_pdf', side_effect=fake_render), \
             patch('core.database.connect_database', side_effect=AssertionError('Retry must not query SQL')), \
             patch('smtplib.SMTP', side_effect=AssertionError('SMTP must not be called')):
            result = execute_bundle(self.cohort, 0, date(2026, 8, 21), 'Tanya', source,
                                    ('BSFin',), ['recipient@example.invalid'], options)
        self.assertEqual(result['status'], 'preview')
        manifest = json.loads((run / 'manifest.json').read_text())
        self.assertEqual(before, manifest['artifacts']['xlsx']['sha256'])
        msg = BytesParser(policy=policy.default).parsebytes((run / 'human.eml').read_bytes())
        self.assertEqual([a.get_content_type() for a in msg.iter_attachments()],
            ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'application/pdf'])
        self.assertIn('BSFin', list(msg.iter_attachments())[1].get_filename())

    def test_no_send_and_uncertain_smtp_are_journaled(self):
        attachment = self.folder / 'report.xlsx'
        attachment.write_bytes(b'fixture')
        manifest = {'routes': {}}
        msg = message('Fixture', 'Body', [attachment], ['test@example.invalid'], message_key='test')
        save_route(self.folder, manifest, 'human', msg, ['test@example.invalid'])
        with patch('smtplib.SMTP', side_effect=AssertionError('no-send called SMTP')):
            deliver(self.folder, manifest, dry_run=True)
        class UncertainSMTP:
            def __init__(self, *args, **kwargs): pass
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def ehlo(self): pass
            def starttls(self, **kwargs): pass
            def login(self, *args): pass
            def send_message(self, *args, **kwargs): raise ConnectionError('lost acknowledgement')
        with patch.dict(os.environ, {'SMTP_PASS': 'mock-only'}):
            with self.assertRaises(ConnectionError):
                deliver(self.folder, manifest, dry_run=False, smtp_factory=UncertainSMTP)
            self.assertEqual(manifest['routes']['human']['state'], 'uncertain')
            with self.assertRaisesRegex(RuntimeError, 'uncertain'):
                deliver(self.folder, manifest, dry_run=False, smtp_factory=UncertainSMTP)


if __name__ == '__main__':
    main()
