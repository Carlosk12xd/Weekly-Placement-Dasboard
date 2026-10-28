"""Build report workbooks for a later Friday with the real builders and mocked SQL.

The supplied fixtures end on 2026-08-21, before the October switch to the Class of 2027.
`build_week` copies a workbook and runs the leadership or program builder for a new date.
Each mocked query answers with the counts in the matching table's newest column, moved
by `shift` students from Actively seeking to Accepted, so every table, the summary and
the population query get consistent, checkable numbers.
"""
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
import re
import shutil
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

POPULATION = [(2027, 'Fall 2026', 'Enrolled', 1500), (2027, 'Spring 2027', 'Graduated', 400), (2026, 'Spring 2026', 'Enrolled', 3)]


def newest_counts(path):
    """{table name: [(status, count)]} from each status table's last column."""
    wb = load_workbook(path)
    try:
        counts = {}
        for ws in wb.worksheets:
            for name in ws.tables:
                left, top, right, bottom = range_boundaries(ws.tables[name].ref)
                if 'search status' not in str(ws.cell(top, left).value or '').lower() or right == left:
                    continue
                rows = []
                for r in range(top + 1, bottom + 1):
                    label = str(ws.cell(r, left).value or '').strip()
                    if label and label.lower() not in ('class size', 'total', '% placed', 'placement %'):
                        value = ws.cell(r, right).value
                        rows.append((label, int(value) if isinstance(value, (int, float)) else 0))
                counts[name] = rows
        return counts
    finally:
        wb.close()


def moved(rows, shift):
    rows = dict(rows)
    step = min(shift, rows.get('Actively seeking', 0))
    if 'Accepted an offer' in rows:
        rows['Accepted an offer'] += step
        rows['Actively seeking'] = rows.get('Actively seeking', 0) - step
    return [(s, n) for s, n in rows.items() if n]


def leadership_answer(counts, shift):
    from LeadershipReport.update_overall_report import PROGRAMS, byprog_full_names, byprog_int_names
    def answer(cur, sql, params=()):
        program = re.search(r"program = '(\w+)'", sql)
        if 'AS offer_accepted' in sql:
            rows = []
            for prog in PROGRAMS:
                c = dict(moved(counts[byprog_full_names(prog)[1]], shift))
                not_seeking = sum(n for s, n in c.items() if s.startswith('Not seeking'))
                rows.append((prog, c.get('Accepted an offer', 0), c.get('Actively seeking', 0), 0,
                             not_seeking, 0, sum(c.values())))
            return rows
        if 'AS grad_term' in sql:
            return POPULATION
        internship = 'internship_search_status' in sql
        if program:
            names = byprog_int_names(program[1]) if internship else byprog_full_names(program[1])
            return moved(counts[names[1]], shift)
        return moved(counts['INT_total_wh' if internship else 'FT_total_wh'], shift)
    return answer


def program_answer(counts, programs, shift):
    from ProgramReports.create_program_reports import table_names
    tables = table_names(programs)
    def answer(cur, sql, params=()):
        program = re.search(r"program = '(\w+)'", sql)
        single_class = re.search(r'WHERE class_of = (\d{4})', sql)
        internship = 'internship_search_status' in sql
        if single_class:  # BSFin's two internship classes, queried next year first
            name = tables['BSFin'][3] if answer.bsfin_calls == 0 else tables['BSFin'][5]
            answer.bsfin_calls += 1
            return moved(counts[name], shift)
        if program:
            names = tables[program[1]]
            return moved(counts[names[3] if internship else names[1]], shift)
        return moved(counts[tables['Class'][3] if internship else tables['Class'][1]], shift)
    answer.bsfin_calls = 0
    return answer


def build_week(source, output, *, cohort, report_date, programs=None, shift=3, capture=None,
               counts_from=None, builder_main=None):
    """Copy `source` to `output` and add `report_date` with the real builder (mocked SQL).
    programs=None builds the leadership report. counts_from: workbook whose newest columns
    answer the SQL (default `source`; pass last year's for a freshly created class workbook).
    builder_main: the unpatched builder, when a test has patched the module's main."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    if Path(source).resolve() != output.resolve():
        shutil.copy2(source, output)
    counts = newest_counts(counts_from or output)
    cursor = SimpleNamespace(close=lambda: None)
    conn = SimpleNamespace(cursor=lambda: cursor, close=lambda: None)
    if programs is None:
        from LeadershipReport import update_overall_report as builder
        with patch.object(builder, 'connect_database', return_value=conn), \
             patch.object(builder, 'fetch_rows', side_effect=leadership_answer(counts, shift)), \
             redirect_stdout(StringIO()):
            (builder_main or builder.main)(cohort.id, cohort.semester_byu, cohort.internship_years, str(output),
                                           report_date=report_date, capture=capture)
    else:
        from ProgramReports import create_program_reports as builder
        with patch.object(builder, 'connect_database', return_value=conn), \
             patch.object(builder, 'fetch_rows', side_effect=program_answer(counts, tuple(programs), shift)), \
             redirect_stdout(StringIO()):
            (builder_main or builder.main)(tuple(programs), cohort, output_path=output,
                                           report_date=report_date, capture=capture)
    return output
