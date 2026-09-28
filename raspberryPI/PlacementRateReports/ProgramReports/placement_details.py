from core.database import connect_database, load_environment
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

from ProgramReports.cd_sql_queries import write_pd_bsfin, write_pd_healthcare, write_pd_program

BASE_DIR = Path(__file__).resolve().parent


LAYOUTS = {
    # default program full-time: start_row, start_col, end_col
    ("DEFAULT", "ft", "default"): (5, 2, 12),
    ("DEFAULT", "int", "default"): (5, 14, 25),

    # healthcare layout uses slightly different columns in your original code (year==8)
    ("Healthcare", "ft", "healthcare"): (5, 2, 13),  
    ("Healthcare", "int", "healthcare"): (5, 15, 26),
    ("Healthcare", "ft", "class_total"): (5, 2, 21),
    ("Healthcare", "int", "class_total"): (5, 2, 21),

    # BSFin special: full-time same as default, but interns are by class_of
    ("BSFin", "ft", "default"): (5, 2, 12),
    ("BSFin", "int", "bsfin_int1"): (5, 15, 25),  # example columns for 2027
    ("BSFin", "int", "bsfin_int2"): (5, 26, 36),  # example columns for 2028
}

# Returns a LIST of TUPLES
def fetch_rows(cur, sql, params=()):
    cur.execute(sql, params)
    return list(cur.fetchall())


def program_to_filename(programs):
    if not programs:
        raise RuntimeError("No programs provided.")
    return programs[0] if len(programs) == 1 else "-".join(programs)

def update_rows(ws, ft_details, layout_params, del_flag):
    start_row, start_col, end_col = layout_params

    if del_flag:
        ws.delete_rows(4, 1000)

    left_side = Side(style="thin")
    right_side = Side(style="thin")

    for i, row in enumerate(ft_details):
        for j, value in enumerate(row):
            col = start_col + j
            cell = ws.cell(row=start_row + i, column = start_col + j, value=value)

            if col == start_col:
                cell.border = Border(left=left_side)
            if col == end_col:
                cell.border = Border(right=right_side)


def get_layout(program: str, section: str, year_key: str = "default"):
    # try exact program-specific layout
    key = (program, section, year_key)
    if key in LAYOUTS:
        return LAYOUTS[key]
    
    raise KeyError(f"No layout found for program={program}, section={section}, year_key={year_key}")

def reset_workbook_view_to_a1(wb):
    """
    Reset each worksheet's visible selection so Excel opens at A1 instead of
    the last edited cell.
    """
    if wb.worksheets:
        wb.active = 0

    for ws in wb.worksheets:
        ws.sheet_view.topLeftCell = "A1"
        if ws.sheet_view.selection:
            ws.sheet_view.selection[0].activeCell = "A1"
            ws.sheet_view.selection[0].sqref = "A1"


# Get SQL Data -> cycle through sheets -> Update both Full and Int tables -> done
def main(programs, cohort, output_path=None, report_date=None):
    conn = connect_database(BASE_DIR)
    cur = conn.cursor()
    try:
        ft_details = {}
        int_details = {}
        total_health_details = {}

        c_year = cohort.id
        int_year1 = int(c_year) + 1
        int_year2 = int(c_year) + 2

        for program in programs:
            if program == 'BSFin':
                queries = write_pd_bsfin(program, cohort)
                ft_details[program] = fetch_rows(cur, queries[0])

                bsfin_interns = {}
            
                bsfin_interns[int_year1] = (fetch_rows(cur, queries[1]))
                bsfin_interns[int_year2] = (fetch_rows(cur, queries[2]))

                int_details[program] = bsfin_interns

            elif program == 'Healthcare':
                queries = write_pd_healthcare(cohort)

                ft_details[program] = fetch_rows(cur, queries[0])
                int_details[program] = fetch_rows(cur, queries[1])

                total_health_details[1] = fetch_rows(cur, queries[2])
                total_health_details[2] = fetch_rows(cur, queries[3])
            else:
                queries = write_pd_program(program, cohort)
                ft_details[program] = fetch_rows(cur, queries[0])
                int_details[program] = fetch_rows(cur, queries[1])

    finally:
        cur.close()
        conn.close()

    FILEPATH_TEMPLATE = str(output_path) if output_path else load_environment(BASE_DIR).get("OUTPUT_PATH", str(BASE_DIR / f"{c_year}_Placement_Reports" / (str(c_year) + "-WeeklyPlacement-{file_label}.xlsx")))
    fileLbl = program_to_filename(programs)
    wb_path = FILEPATH_TEMPLATE.format(file_label=fileLbl)
    wb = load_workbook(wb_path, data_only=False)

    for program in programs:
        print(f"Starting {program}...")
        if program == 'BSFin':
            ws_name = f'{program} Placement Details'
            ws = wb[ws_name]

            ft_layout = get_layout(program, "ft", "default")
            update_rows(ws, ft_details[program], ft_layout, del_flag=True)

            int_2027_layout = get_layout("BSFin", "int", "bsfin_int1")
            update_rows(ws, int_details[program][int_year1], int_2027_layout, del_flag=False)

            int_2028_layout = get_layout("BSFin", "int", "bsfin_int2")            
            update_rows(ws, int_details[program][int_year2], int_2028_layout, del_flag=False)

        elif program == 'Healthcare':
            ws1 = "Healthcare Full-Time"
            ws2 = "Healthcare Internships"
            ws3 = "Past Week Placement Details"

            total_ft_layout = get_layout(program, "ft", "class_total")
            update_rows(wb[ws1], total_health_details[1], total_ft_layout, del_flag=True)

            total_int_layout = get_layout(program, "int", "class_total")
            update_rows(wb[ws2], total_health_details[2], total_int_layout, del_flag=True)

            weekly_ft_layout = get_layout(program, "ft", "healthcare")
            weekly_int_layout = get_layout(program, "int", "healthcare")

            update_rows(wb[ws3], ft_details[program], weekly_ft_layout, del_flag=True)
            update_rows(wb[ws3], int_details[program], weekly_int_layout, del_flag=False)
            
        else:
            ws_name = f'{program} Placement Details'
            ws = wb[ws_name]

            ft_layout = get_layout("DEFAULT", "ft", "default")
            int_layout = get_layout("DEFAULT", "int", "default")

            update_rows(ws, ft_details[program], ft_layout, del_flag=True)
            update_rows(ws, int_details[program], int_layout, del_flag=False)

    reset_workbook_view_to_a1(wb)

    wb.save(wb_path)
    print(f"Updated: {wb_path}")
    return str(wb_path)


if __name__ == "__main__":
    # example
    main(["MPA"])
