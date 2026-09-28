"""Narrow OOXML edits that preserve the supplied dashboard's chart/style parts."""
from pathlib import Path
import copy
import posixpath
import re
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as ET
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'s': S, 'c': C, 'a': A}


class Package:
    def __init__(self, path):
        with ZipFile(path) as archive:
            self.parts = {n: archive.read(n) for n in archive.namelist()}
        self.trees = {}
        self.strings = []
        if 'xl/sharedStrings.xml' in self.parts:
            self.strings = [''.join(x.itertext()) for x in self.xml('xl/sharedStrings.xml')]
        wb = self.xml('xl/workbook.xml')
        relations = self.relationships('xl/workbook.xml')
        self.sheets = {s.get('name'): relations[s.get(f'{{{R}}}id')]
                       for s in wb.find(f'{{{S}}}sheets')}

    def xml(self, part):
        if part not in self.trees:
            self.trees[part] = ET.fromstring(self.parts[part])
        return self.trees[part]

    def relationships(self, part):
        folder, name = posixpath.split(part)
        relpath = f'{folder}/_rels/{name}.rels'
        if relpath not in self.parts:
            return {}
        return {x.get('Id'): (x.get('Target').lstrip('/') if x.get('Target').startswith('/')
                 else posixpath.normpath(posixpath.join(folder, x.get('Target'))))
                for x in self.xml(relpath) if x.get('TargetMode') != 'External'}

    def cell(self, sheet, address, create=False, style_from=None):
        root = self.xml(self.sheets[sheet])
        data = root.find(f'{{{S}}}sheetData')
        rn, cn = coordinate_to_tuple(address.replace('$', ''))
        row = data.find(f"s:row[@r='{rn}']", NS)
        if row is None:
            if not create:
                return None
            row = ET.Element(f'{{{S}}}row', r=str(rn))
            following = next((r for r in data if int(r.get('r')) > rn), None)
            data.insert(data.index(following) if following is not None else len(data), row)
        cell = row.find(f"s:c[@r='{address}']", NS)
        if cell is None and create:
            cell = ET.Element(f'{{{S}}}c', r=address)
            following = next((c for c in row if coordinate_to_tuple(c.get('r'))[1] > cn), None)
            row.insert(row.index(following) if following is not None else len(row), cell)
            if style_from:
                source = self.cell(sheet, style_from)
                if source is not None and source.get('s') is not None:
                    cell.set('s', source.get('s'))
        return cell

    def get(self, sheet, address):
        cell = self.cell(sheet, address.replace('$', ''))
        if cell is None:
            return None
        if cell.get('t') == 'inlineStr':
            return ''.join(cell.find(f'{{{S}}}is').itertext())
        value = cell.find(f'{{{S}}}v')
        if value is None or value.text is None:
            return None
        if cell.get('t') == 's':
            return self.strings[int(value.text)]
        if cell.get('t') in ('str', 'e'):
            return value.text
        return float(value.text)

    def set(self, sheet, address, value, style_from=None):
        cell = self.cell(sheet, address, True, style_from)
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop('t', None)
        if value is None:
            return
        if isinstance(value, str):
            cell.set('t', 'inlineStr')
            ET.SubElement(ET.SubElement(cell, f'{{{S}}}is'), f'{{{S}}}t').text = value
        else:
            ET.SubElement(cell, f'{{{S}}}v').text = str(value)

    def tables(self, sheet):
        rels = self.relationships(self.sheets[sheet])
        root = self.xml(self.sheets[sheet])
        return [self.xml(rels[x.get(f'{{{R}}}id')])
                for x in root.findall('s:tableParts/s:tablePart', NS)]

    def resize_table(self, sheet, table, last_column):
        start, end = table.get('ref').split(':')
        top, _ = coordinate_to_tuple(start)
        bottom, _ = coordinate_to_tuple(end)
        table.set('ref', f'A{top}:{get_column_letter(last_column)}{bottom}')
        af = table.find(f'{{{S}}}autoFilter')
        if af is not None:
            af.set('ref', table.get('ref'))
        cols = table.find(f'{{{S}}}tableColumns')
        proto = copy.deepcopy(cols[-1])
        for child in list(cols):
            cols.remove(child)
        for col in range(1, last_column + 1):
            child = copy.deepcopy(proto)
            child.set('id', str(col))
            child.set('name', str(self.get(sheet, f'{get_column_letter(col)}{top}')))
            for item in list(child):
                child.remove(item)
            cols.append(child)
        cols.set('count', str(last_column))

    def chart_parts(self):
        return [p for p in self.parts if re.fullmatch(r'xl/charts/chart\d+\.xml', p)]

    def remove_calc_chain(self):
        self.parts.pop('xl/calcChain.xml', None)
        self.trees.pop('xl/calcChain.xml', None)
        for name in ['xl/_rels/workbook.xml.rels', '[Content_Types].xml']:
            root = self.xml(name)
            for child in list(root):
                if 'calcChain' in str(child.attrib):
                    root.remove(child)
        root = self.xml('xl/workbook.xml')
        calc = root.find(f'{{{S}}}calcPr')
        if calc is None:
            calc = ET.SubElement(root, f'{{{S}}}calcPr')
        calc.set('fullCalcOnLoad', '1')

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + '.tmp')
        with ZipFile(temporary, 'w', ZIP_DEFLATED) as archive:
            for name, payload in self.parts.items():
                if name in self.trees:
                    payload = ET.tostring(self.trees[name], xml_declaration=True, encoding='UTF-8', standalone=True)
                archive.writestr(name, payload)
        temporary.replace(path)

